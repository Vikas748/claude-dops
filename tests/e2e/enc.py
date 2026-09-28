exec(open('/tmp/up.py').read().split("# ---------- setup ----------")[0])
admin=web_login("head.dept@hospital.in")
call("POST","/api/admin",{"name":"Nurse Lep","email":"nurse.lep@hospital.in","role":"NURSE","status":"ACTIVE","permissions":["LEPROSY:VIEW","LEPROSY:EDIT","LEPROSY:CREATE"]},cookie=admin)
call("POST","/api/admin",{"name":"Skin Only","email":"skin@hospital.in","role":"STAFF","status":"ACTIVE","permissions":["SKIN_BANK:VIEW"]},cookie=admin)
P={"Age":"45","Sex":"Male","Address":"कटनी","Mobile No.":"9876543210","Aadhaar Card No.":"123412341234","Samagra ID":"S1","Ayushman Card":"A1","Diagnosis":"Claw hand","Date of Admission":"2026-09-01","Date of Surgery":"2026-09-03","Bank Account No.":"000123456789","Bank Name":"SBI","Amount Released":"No","Amount":"12000"}
c,d,_=call("POST","/api/special",{"kind":"LEPROSY","recordDate":"2026-09-05","primaryName":"रामप्रसाद","payload":P},cookie=admin)
raw_db=sql("select payload from special_records where kind='LEPROSY'")
ok(c in (200,201) and "123412341234" not in raw_db and "000123456789" not in raw_db and raw_db.count("enc:v1:")==2, f"E1 saved; database holds ciphertext only for Aadhaar + bank account (plaintext absent)")
rid=int(sql("select id from special_records where kind='LEPROSY'"))
Q="/api/special?kind=LEPROSY&from=2026-09-01&to=2026-09-30"
_,d,_=call("GET",Q,cookie=admin); pa=d["data"][0]["payload"]
ok(pa["Aadhaar Card No."]=="123412341234" and pa["Bank Account No."]=="000123456789", "E2 admin sees full numbers (decrypted)")
nurse=web_login("nurse.lep@hospital.in"); _,d,_=call("GET",Q,cookie=nurse); pn=d["data"][0]["payload"]
ok(pn["Aadhaar Card No."]=="XXXXXXXX1234" and pn["Bank Account No."]=="XXXXXXXX6789", f"E3 nurse sees masked: {pn['Aadhaar Card No.']}, {pn['Bank Account No.']}")
pn["Diagnosis"]="Claw hand – post op"
c,d,_=call("POST","/api/special",{"id":rid,"kind":"LEPROSY","recordDate":"2026-09-05","primaryName":"रामप्रसाद","payload":pn},cookie=nurse)
_,d2,_=call("GET",Q,cookie=admin); pa=d2["data"][0]["payload"]
ok(c==200 and pa["Aadhaar Card No."]=="123412341234" and pa["Bank Account No."]=="000123456789" and pa["Diagnosis"].endswith("post op"),
   f"E4 nurse edit with masked values -> {c}; real numbers kept, diagnosis updated (before the fix: Aadhaar error / bank number overwritten)")
skin=web_login("skin@hospital.in"); c,d,_=call("GET",f"/api/special/columns?historyId={rid}",cookie=skin)
ok(c==403, f"E5 Skin-Bank-only user asks for Leprosy history -> {c} (previously returned unmasked data)")
c,d,_=call("GET",f"/api/special/columns?historyId={rid}",cookie=nurse); h=d["data"][0]["payload"]
ok(c==200 and h["Aadhaar Card No."]=="XXXXXXXX1234", f"E6 nurse history -> {c}, masked")
c,d,_=call("GET",f"/api/special/columns?historyId={rid}",cookie=admin); ok(d["data"][0]["payload"]["Aadhaar Card No."]=="123412341234", "E7 admin history decrypted")
s,b,_=raw("GET",B+Q+"&format=xlsx",headers={"cookie":admin}); open("/tmp/enc.xlsx","wb").write(b)
import openpyxl; ws=openpyxl.load_workbook("/tmp/enc.xlsx").active; hdr=[c.value for c in ws[7]]; row=[c.value for c in ws[8]]
ok(row[hdr.index("Aadhaar Card No.")]=="123412341234" and row[hdr.index("Bank Account No.")]=="000123456789", "E8 admin Excel export has real numbers")
# legacy plaintext row is readable and gets encrypted on next save
sql("""insert into special_records (kind,record_date,primary_name,status,payload,created_at,updated_at) values ('LEPROSY','2026-09-06','Legacy','ACTIVE','{\\"Aadhaar Card No.\\":\\"999988887777\\",\\"Bank Account No.\\":\\"111\\"}',now()::text,now()::text)""")
lid=int(sql("select id from special_records where primary_name='Legacy'"))
_,d,_=call("GET",Q,cookie=admin); leg=[x for x in d["data"] if x["primaryName"]=="Legacy"][0]["payload"]
ok(leg["Aadhaar Card No."]=="999988887777", "E9 old plaintext record still readable")
L={**P,"Aadhaar Card No.":"999988887777","Bank Account No.":"111"}
call("POST","/api/special",{"id":lid,"kind":"LEPROSY","recordDate":"2026-09-06","primaryName":"Legacy","payload":L},cookie=admin)
ok("999988887777" not in sql(f"select payload from special_records where id={lid}"), "E10 ...and encrypted on its next save")
# tamper: copy the Aadhaar ciphertext into the bank field -> refused (field name is authenticated)
sql(f"update special_records set payload = jsonb_set(payload::jsonb, '{{Bank Account No.}}', payload::jsonb->'Aadhaar Card No.')::text where id={rid}")
_,d,_=call("GET",Q,cookie=admin); t=[x for x in d["data"] if x["id"]==rid][0]["payload"]
ok(t["Bank Account No."]=="[unreadable]" and t["Aadhaar Card No."]=="123412341234", f"E11 ciphertext moved to another field is rejected ({t['Bank Account No.']}), list still loads")
